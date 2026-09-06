# -*- coding: utf-8 -*-
"""
Read Departments_Dashboard_Linked.xlsx and produce a compact JSON payload
with the same aggregates the Excel Dashboard sheet computes (latest-day
breakdowns, latest KPIs, daily trend). This JSON gets baked into the
static Dashboard website HTML at publish time.
"""
import json
import sys
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else "Departments_Dashboard_Linked.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "dashboard_data.json"

wb = openpyxl.load_workbook(SRC, data_only=False)


def header_index(ws, header_row, name):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if isinstance(v, str) and v.strip() == name:
            return c
    return None


def extract_rows(sheet_name, colmap, anchor_key, header_row=1):
    """colmap: {friendly_name: exact_header_text}. anchor_key: friendly
    name whose column must hold a number for the row to count as data
    (skips blank rows and the LEGEND row)."""
    ws = wb[sheet_name]
    idx = {k: header_index(ws, header_row, h) for k, h in colmap.items()}
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        anchor_val = ws.cell(row=r, column=idx[anchor_key]).value
        if not isinstance(anchor_val, (int, float)):
            continue
        rec = {}
        for k, c in idx.items():
            rec[k] = ws.cell(row=r, column=c).value
        out.append(rec)
    return out


credit_decisions = extract_rows(
    "Credit_Decisions",
    {
        "date": "Date - التاريخ",
        "type": "Credit Decision - قرار الائتمان",
        "count": "Number of Requests - عدد الطلبات",
        "amount": "Amounts - المبالغ",
    },
    anchor_key="count",
)

credit_kpis = extract_rows(
    "Credit_KPIs",
    {
        "date": "Date - التاريخ",
        "today_financing": "Today Financing Amount",
        "total_financing": "Total Financing Amount",
        "total_contracts": "Total Number Contracts",
        "cancelled": "Cancelled",
    },
    anchor_key="total_financing",
)

sales = extract_rows(
    "Sales",
    {
        "date": "Date - التاريخ",
        "status": "Status - الحالة",
        "count": "Count - العدد",
        "amount": "Amounts - المبالغ",
    },
    anchor_key="count",
)

customer_care = extract_rows(
    "Customer_Care",
    {
        "date": "Date - التاريخ",
        "parameter": "Parameter - المعيار",
        "active": "Active - نشط",
        "close": "Close - مغلق",
    },
    anchor_key="active",
)

collection_summary = extract_rows(
    "Collection_Summary",
    {
        "date_from": "Date From - من",
        "date_to": "Date To - الى",
        "item": "Item - البند",
        "amount": "Amount - المبلغ",
    },
    anchor_key="amount",
)

# Fixed_Arrears: header row is row 3; description column excludes the
# total row (its Description cell is blank).
fixed_arrears_rows = extract_rows(
    "Fixed_Arrears",
    {
        "bucket": "Bucket - البوكت",
        "description": "Description - الوصف",
        "customers": "Number of Customers - عدد العملاء",
        "amount": "Overdue Amount - مبلغ المتأخرات",
    },
    anchor_key="amount",
    header_row=3,
)
fixed_arrears_rows = [r for r in fixed_arrears_rows if r.get("description")]
fixed_arrears_asof = wb["Fixed_Arrears"]["B1"].value


def to_str_date(v):
    return str(v).strip() if v is not None else None


for rows, key in [
    (credit_decisions, "date"), (credit_kpis, "date"), (sales, "date"),
    (customer_care, "date"), (collection_summary, "date_from"),
]:
    for r in rows:
        if key in r:
            r[key] = to_str_date(r[key])
for r in collection_summary:
    r["date_to"] = to_str_date(r["date_to"])


def max_date(rows, key):
    vals = [r[key] for r in rows if r.get(key)]
    return max(vals) if vals else None


def group_sum(rows, group_key, sum_keys):
    out = {}
    for r in rows:
        g = r[group_key]
        if g not in out:
            out[g] = {k: 0 for k in sum_keys}
        for k in sum_keys:
            out[g][k] += r.get(k) or 0
    return out


# ---- Latest-day breakdowns -------------------------------------------------
cd_latest_date = max_date(credit_decisions, "date")
cd_latest = [r for r in credit_decisions if r["date"] == cd_latest_date]

sales_latest_date = max_date(sales, "date")
sales_latest = [r for r in sales if r["date"] == sales_latest_date]

care_latest_date = max_date(customer_care, "date")
care_latest = [r for r in customer_care if r["date"] == care_latest_date]

coll_latest_period = max_date(collection_summary, "date_to")
coll_latest = [r for r in collection_summary if r["date_to"] == coll_latest_period]

kpi_latest_date = max_date(credit_kpis, "date")
kpi_latest = next((r for r in credit_kpis if r["date"] == kpi_latest_date), None)

total_collections = sum(r["amount"] for r in coll_latest if r["item"] == "الاجمالي")
fixed_arrears_total = sum(r["amount"] for r in fixed_arrears_rows)
sales_accepted_latest = sum(
    r["count"] for r in sales_latest if r["status"] == "Accepted - تم التنفيذ"
)
care_active_latest = sum(r["active"] for r in care_latest)

# ---- Trend series (all dates present) -------------------------------------
all_dates = sorted(set(
    [r["date"] for r in credit_kpis if r.get("date")]
    + [r["date_to"] for r in collection_summary if r.get("date_to")]
))
kpi_by_date = {r["date"]: r for r in credit_kpis}
coll_total_by_date = {
    r["date_to"]: r["amount"] for r in collection_summary if r["item"] == "الاجمالي"
}
trend = {
    "dates": all_dates,
    "total_financing": [kpi_by_date[d]["total_financing"] if d in kpi_by_date else None for d in all_dates],
    "total_collections": [coll_total_by_date.get(d) for d in all_dates],
}

payload = {
    "generated_note": "قيم حقيقية من ملف Excel المُرسل — تتحدث عند إعادة نشر الموقع بملف جديد.",
    "kpi": {
        "as_of_date": kpi_latest_date,
        "total_financing": (kpi_latest or {}).get("total_financing", 0),
        "today_financing": (kpi_latest or {}).get("today_financing", 0),
        "total_contracts": (kpi_latest or {}).get("total_contracts", 0),
        "cancelled": (kpi_latest or {}).get("cancelled", 0),
        "total_collections": total_collections,
        "collections_period_to": coll_latest_period,
        "fixed_arrears_total": fixed_arrears_total,
        "fixed_arrears_asof": to_str_date(fixed_arrears_asof),
        "sales_accepted_latest": sales_accepted_latest,
        "sales_date": sales_latest_date,
        "care_active_latest": care_active_latest,
        "care_date": care_latest_date,
    },
    "breakdowns": {
        "credit_decisions": {
            "date": cd_latest_date,
            "labels": [r["type"] for r in cd_latest],
            "counts": [r["count"] for r in cd_latest],
            "amounts": [r["amount"] for r in cd_latest],
        },
        "sales": {
            "date": sales_latest_date,
            "labels": [r["status"] for r in sales_latest],
            "counts": [r["count"] for r in sales_latest],
        },
        "customer_care": {
            "date": care_latest_date,
            "labels": [r["parameter"] for r in care_latest],
            "active": [r["active"] for r in care_latest],
            "close": [r["close"] for r in care_latest],
        },
        "collections": {
            "period_to": coll_latest_period,
            "labels": [r["item"] for r in coll_latest if r["item"] != "الاجمالي"],
            "amounts": [r["amount"] for r in coll_latest if r["item"] != "الاجمالي"],
        },
        "fixed_arrears": {
            "as_of": to_str_date(fixed_arrears_asof),
            "labels": [str(r["bucket"]) for r in fixed_arrears_rows],
            "amounts": [r["amount"] for r in fixed_arrears_rows],
            "customers": [r["customers"] for r in fixed_arrears_rows],
        },
    },
    "trend": trend,
}

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(payload, f, ensure_ascii=False, indent=2)

print("Wrote", OUT)
print(json.dumps(payload, ensure_ascii=False, indent=2)[:2000])
